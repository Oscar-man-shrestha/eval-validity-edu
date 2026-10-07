#!/usr/bin/env python3
"""Build IEEE A4 conference Word draft for advisor / pre-journal review.

Confirmatory H-rev cells stay pending_confirmatory until OSF + --touch-test.
Exploratory numbers load from outputs_junyi/phases/*.json.
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
PHASES = ROOT / "outputs_junyi" / "phases"
OUT = Path(__file__).resolve().parent / "NeuroTrace-DAG_IEEE_A4_Conference.docx"


def load(name: str) -> dict:
    p = PHASES / name
    return json.loads(p.read_text()) if p.exists() else {}


def set_run(run, *, size=10, bold=False, italic=False, font="Times New Roman"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(doc, text, *, size=10, bold=False, italic=False, center=False, space_after=6, space_before=0, first_indent=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if first_indent is not None:
        pf.first_line_indent = Cm(first_indent)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, italic=italic)
    return p


def add_heading_ieee(doc, text, level=1):
    # IEEE: Heading 1 often ALL CAPS
    label = text.upper() if level == 1 else text
    size = 10 if level == 1 else 10
    p = add_para(doc, label, size=size, bold=True, space_before=10, space_after=4)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        set_run(r, size=8, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            set_run(r, size=8)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if col_widths:
        for row in table.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = Cm(w)
    doc.add_paragraph()
    return table


def set_two_columns(section):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "708")  # ~0.5"


def build():
    rr = load("review_ranking.json")
    full = load("paper_ranking_full.json")
    p1 = load("phase1_prepare.json")
    p2 = load("phase2_embeddings.json")
    p3 = load("phase3_memory.json")
    cal = load("paper_calibration.json")
    m = rr.get("metrics") or {}
    ragr = m.get("ragr_r5") or {}
    recent = m.get("recent5_r5") or {}
    full_ragr = (full.get("metrics") or {}).get("ragr_r5") or {}
    mem = (p3.get("evaluate_on_test_json") or {}).get("logistic_step_dt") or {}
    gate = (cal.get("gate") or {})
    pref = (cal.get("p_recall_on_continue_queries") or {})

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.57)
    section.right_margin = Cm(1.57)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(2.54)

    # Title (IEEE: no symbols in title ideally — avoid special arrows)
    add_para(
        doc,
        "When Next-Item Metrics Mislead: Evaluation Validity Across Educational Interaction Logs",
        size=24,
        bold=True,
        center=True,
        space_after=12,
    )
    add_para(doc, "Oscar Man Shrestha", size=11, center=True, space_after=0)
    add_para(doc, "Dept. name of organization (Affiliation)", size=9, italic=True, center=True, space_after=0)
    add_para(doc, "City, Country", size=9, italic=True, center=True, space_after=0)
    add_para(doc, "email address or ORCID", size=9, italic=True, center=True, space_after=10)

    # Abstract — IEEE uses Abstract— prefix, no math symbols if avoidable
    abs_p = doc.add_paragraph()
    abs_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abs_p.paragraph_format.space_after = Pt(6)
    r0 = abs_p.add_run("Abstract—")
    set_run(r0, size=9, bold=True, italic=True)
    abs_body = (
        "Recall@K is the usual report card for next-item models on student logs. "
        "We are less sure it measures what people think it measures. "
        f"On Junyi Academy ({p1.get('concepts', 835)} concepts; "
        f"{p1.get('train_sequences', 27434):,} / {p1.get('test_sequences', 6290):,} train/test sequences), "
        f"a gated ranker gets Recall@5 {ragr.get('mean', 0.891):.3f} "
        f"[{ragr.get('ci', [0.880, 0.902])[0]:.3f}, {ragr.get('ci', [0.880, 0.902])[1]:.3f}], "
        f"yet Recent-5 is already {recent.get('mean', 0.885):.3f}—within a 0.01 practical gap—so we do not claim a recommender win. "
        "The more interesting question is whether the action unit (question, skill, or topic cluster) "
        "changes which simple model looks better. "
        "We lock a small confirmatory comparison of a GRU probe against first-order Markov on Junyi timed, "
        "ASSISTments 2009, and XES3G5M, at native units and at matched clusters (~120), under a freeze holdout "
        "and Bonferroni family m_planned=6. "
        "Holdout cells stay blank until external registration and one touch-test; "
        "Junyi ranking numbers below are exploratory."
    )
    r1 = abs_p.add_run(abs_body)
    set_run(r1, size=9, italic=True)

    kw = doc.add_paragraph()
    kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rk = kw.add_run("Keywords—")
    set_run(rk, size=9, bold=True, italic=True)
    rk2 = kw.add_run(
        "educational data mining, evaluation validity, next-item recommendation, "
        "knowledge tracing logs, preregistration, action granularity"
    )
    set_run(rk2, size=9, italic=True)

    # Switch to two columns after front matter
    # python-docx: continuous section break then two cols is fiddly; keep single column
    # for advisor Word readability (IEEE camera-ready can reflow later).

    add_heading_ieee(doc, "I. Introduction")
    add_para(
        doc,
        "On educational logs, next-item Recall@K often rewards repeats and the curriculum’s natural "
        "order. That is not the same as measuring whether a model captured learning structure. Worse, "
        "changing how we define an “item”—question id versus skill or topic cluster—can flip which "
        "baseline wins. Our goal here is modest: check whether that unit choice systematically "
        "reverses a GRU-versus-Markov contrast under a locked protocol. We are not proposing a new "
        "recommender champion.",
        first_indent=0.5,
    )
    add_para(
        doc,
        "Registered question (prereg-v1): does the sign of the paired GRU-minus-Markov non-repeat "
        "Recall@5 contrast flip between the native unit and the registered cluster unit on at least "
        "one dataset? Until registration and a single freeze holdout run, those cells stay blank "
        "(pending). The Junyi ranking and cross-dataset grids below are exploratory; they are here "
        "so the system is inspectable before the holdout is touched.",
        first_indent=0.5,
    )

    add_heading_ieee(doc, "II. Related Work")
    add_para(
        doc,
        "Knowledge tracing and educational recommender papers usually score platform-native ids. "
        "Duplicates, unit mismatch, and curriculum order show up often as caveats; here they are the "
        "object of study. Library-faithful SASRec or RecBole GRU4Rec are outside this confirmatory "
        "packet for now; the probes below are diagnostic [1]–[3].",
        first_indent=0.5,
    )

    add_heading_ieee(doc, "III. Datasets and Units")
    add_para(
        doc,
        f"Junyi (primary product log): {p1.get('concepts', 835)} English concept names, "
        f"{p1.get('edges', 978)} prerequisite edges, train interactions "
        f"{p1.get('interactions_train', 2342784):,}. "
        f"Embedding sanity: pair-AUC versus random {p2.get('pair_auc_vs_random', 0.904):.3f}; "
        f"same-topic control {p2.get('pair_auc_vs_same_topic', 0.464):.3f}. "
        "Cross-dataset units: Junyi topics, ASSISTments composite skill tokens, XES3G5M questions. "
        "Registered cluster maps target about 120 groups (text MiniLM k-means primary on Junyi and "
        "ASSISTments; co-occurrence secondary for a two-method intersection rule; XES co-occurrence only).",
        first_indent=0.5,
    )
    add_para(doc, "TABLE I. REGISTERED ACTION UNITS", size=8, bold=True, center=True, space_before=8, space_after=2)
    add_table(
        doc,
        ["Dataset", "Native unit", "Cluster unit (registered)"],
        [
            ["Junyi timed", "topic / exercise id", "text k-means (+ cooc intersection)"],
            ["ASSISTments 2009", "composite skill token", "text on skill_name (+ cooc)"],
            ["XES3G5M", "question id", "co-occurrence only (amendment)"],
        ],
        col_widths=[4.0, 5.0, 6.5],
    )

    add_heading_ieee(doc, "IV. Methods")
    add_para(
        doc,
        "A. Confirmatory comparison",
        bold=True,
        space_before=4,
        space_after=2,
    )
    add_para(
        doc,
        "Learners are split once (seed 20261004). Confirmatory GRU probes retrain on freeze train+val "
        "with the frozen winner config and a fixed best_epoch—no peeking at test for early stopping— "
        "over seeds 20261101–20261103. Markov is just a first-order successor table. We planned six "
        "family members (three datasets × native/cluster) and use Bonferroni at 1−0.05/m; ineligible "
        "cells drop out. A cell needs at least 500 non-repeat queries; CIs use ≥10,000 sequence "
        "bootstraps. A dataset supports the claim only if native and cluster both come back "
        "supported_* with opposite signs; the project needs that on at least one dataset. Junyi and "
        "ASSISTments cluster cells require text and co-occurrence maps to agree. Models in the "
        "confirmatory tables: popularity, recency, markov_order1, gru_probe.",
        first_indent=0.5,
    )
    add_para(
        doc,
        "B. Exploratory Junyi work",
        bold=True,
        space_before=4,
        space_after=2,
    )
    add_para(
        doc,
        "Phases 1–6 build embeddings, revisit-memory classifiers, a prerequisite DAG, a review/advance "
        "gate, and a gated ranker (RAGR). Useful context—not part of the confirmatory family.",
        first_indent=0.5,
    )

    add_heading_ieee(doc, "V. Results")
    add_para(
        doc,
        "A. Primary confirmatory cells (pending)",
        bold=True,
        space_before=4,
        space_after=2,
    )
    add_para(
        doc,
        "Tag prereg-v1 freezes the protocol and artifacts. We leave the freeze-holdout scores blank "
        "until OSF/Zenodo registration and one --touch-test—that is deliberate, not incomplete "
        "drafting. Table II therefore says pending in every cell. Validation dual-map contrasts from "
        "before registration are disclosure only; please do not read them as the confirmatory answer.",
        first_indent=0.5,
    )
    add_para(doc, "TABLE II. CONFIRMATORY UNIT CELLS (PENDING TOUCH-TEST)", size=8, bold=True, center=True, space_before=8, space_after=2)
    add_table(
        doc,
        ["Dataset", "Unit", "Verdict", "Bonferroni CI"],
        [
            ["junyi_timed", "native", "pending", "—"],
            ["junyi_timed", "cluster", "pending", "—"],
            ["assistments", "native", "pending", "—"],
            ["assistments", "cluster", "pending", "—"],
            ["xes3g5m", "native", "pending", "—"],
            ["xes3g5m", "cluster", "pending", "—"],
        ],
        col_widths=[3.5, 2.5, 5.0, 3.5],
    )

    add_para(
        doc,
        "B. Exploratory Junyi ranking and calibration",
        bold=True,
        space_before=4,
        space_after=2,
    )
    add_para(
        doc,
        f"Canonical capped evaluation (review_ranking.json, n_seq={ragr.get('n_sequences', 800)}): "
        f"RAGR Recall@5 = {ragr.get('mean', 0.891):.3f} "
        f"[{ragr.get('ci', [0.880, 0.902])[0]:.3f}, {ragr.get('ci', [0.880, 0.902])[1]:.3f}]; "
        f"Recent-5 = {recent.get('mean', 0.885):.3f} "
        f"[{recent.get('ci', [0.872, 0.897])[0]:.3f}, {recent.get('ci', [0.872, 0.897])[1]:.3f}]. "
        f"Full-test RAGR Recall@5 = {full_ragr.get('mean', 0.889):.3f} "
        f"[{full_ragr.get('ci', [0.885, 0.893])[0]:.3f}, {full_ragr.get('ci', [0.885, 0.893])[1]:.3f}] "
        f"(n_seq={full_ragr.get('n_sequences', 6290)}). "
        f"Revisit-memory logistic AUC on official test revisit rows = {mem.get('auc', 0.864):.3f}. "
        f"Gate calibration (full-test queries): Brier={gate.get('brier', 0.0928):.4f}, "
        f"ECE={gate.get('ece', 0.0335):.4f}. "
        "The sub-0.01 RAGR-versus-Recent-5 gap is why we do not sell this as a ranking breakthrough.",
        first_indent=0.5,
    )
    add_para(doc, "TABLE III. EXPLORATORY JUNYI RANKING (NOT CONFIRMATORY)", size=8, bold=True, center=True, space_before=8, space_after=2)
    add_table(
        doc,
        ["Source", "RAGR R@5", "95% CI", "n seq"],
        [
            [
                "review_ranking (capped)",
                f"{ragr.get('mean', 0.891):.3f}",
                f"[{ragr.get('ci', [0.880, 0.902])[0]:.3f}, {ragr.get('ci', [0.880, 0.902])[1]:.3f}]",
                str(ragr.get("n_sequences", 800)),
            ],
            [
                "paper_ranking_full",
                f"{full_ragr.get('mean', 0.889):.3f}",
                f"[{full_ragr.get('ci', [0.885, 0.893])[0]:.3f}, {full_ragr.get('ci', [0.885, 0.893])[1]:.3f}]",
                str(full_ragr.get("n_sequences", 6290)),
            ],
            [
                "Recent-5 (capped)",
                f"{recent.get('mean', 0.885):.3f}",
                f"[{recent.get('ci', [0.872, 0.897])[0]:.3f}, {recent.get('ci', [0.872, 0.897])[1]:.3f}]",
                str(recent.get("n_sequences", 800)),
            ],
        ],
        col_widths=[5.0, 2.5, 4.5, 2.5],
    )

    add_para(
        doc,
        "C. Cross-dataset advance slice (exploratory / diagnostic)",
        bold=True,
        space_before=4,
        space_after=2,
    )
    add_para(
        doc,
        "On advance non-repeat slices, first-order Markov looks strong on curriculum-heavy XES "
        "questions, while GRU probes move around by dataset and unit. That pattern is why we "
        "preregistered the sign-flip test; the grid itself is not the confirmatory answer.",
        first_indent=0.5,
    )
    add_para(doc, "TABLE IV. ADVANCE R@5 DIAGNOSTIC GRID (EXPLORATORY)", size=8, bold=True, center=True, space_before=8, space_after=2)
    add_table(
        doc,
        ["Method", "Junyi timed", "ASSISTments", "XES3G5M"],
        [
            ["popularity", "0.157", "0.214", "0.011"],
            ["markov_order1", "0.417", "0.510", "0.727"],
            ["gru_probe", "0.397", "0.552", "0.612"],
            ["gru4rec_style", "0.407", "0.654", "0.689"],
            ["attn_probe", "0.259", "0.430", "0.589"],
        ],
        col_widths=[4.0, 3.5, 3.5, 3.5],
    )

    add_heading_ieee(doc, "VI. Discussion")
    add_para(
        doc,
        "What we registered is whether changing the action unit reverses the GRU–Markov contrast on "
        "the freeze holdout—not whether a gated ranker beats Recent-k on Junyi. A null or mixed "
        "result still counts: it would mean the flip did not show up under the locked rules. Native "
        "XES stays a descriptive curriculum-heavy case even though it sits in the family under an "
        "explicit amendment.",
        first_indent=0.5,
    )

    add_heading_ieee(doc, "VII. Limitations")
    add_para(
        doc,
        "We do not yet have a library-faithful SASRec/RecBole baseline in confirmatory scope. "
        "ASSISTments and XES redistribution licences are still unverified. Text and co-occurrence "
        "cluster maps disagree (low ARI); intersection is deliberately conservative. And of course "
        "there are no confirmatory holdout numbers until registration and one touch-test.",
        first_indent=0.5,
    )

    add_heading_ieee(doc, "VIII. Conclusion")
    add_para(
        doc,
        "We locked a small confirmatory test about action units and left the Junyi product numbers "
        "clearly exploratory. After registration, one touch-test fills Table II from primary_claims "
        "only; until then those cells stay blank on purpose.",
        first_indent=0.5,
    )

    add_heading_ieee(doc, "Acknowledgment")
    add_para(
        doc,
        "Thanks to advisors and collaborators for comments on framing and scope. "
        "Code and prereg packet: https://github.com/Oscar-man-shrestha/eval-validity-edu (tag prereg-v1).",
        first_indent=0.5,
    )

    add_heading_ieee(doc, "References")
    refs = [
        '[1] B. Hidasi, A. Karatzoglou, L. Baltrunas, and D. Tikk, “Session-based recommendations with recurrent neural networks,” in Proc. ICLR, 2016.',
        '[2] W.-C. Kang and J. McAuley, “Self-attentive sequential recommendation,” in Proc. IEEE ICDM, 2018, pp. 197–206.',
        '[3] F. Wang et al., “EduData: A unified library for educational data mining,” GitHub repository, accessed 2026.',
        '[4] X. Xiong, Z. Zhao, E. G. Van Inwegen, and J. E. Beck, “Going deeper with deep knowledge tracing,” in Proc. EDM, 2016.',
        '[5] P. I. Pavlik Jr., H. Cen, and K. R. Koedinger, “Performance factors analysis—A new alternative to knowledge tracing,” in Proc. AIED, 2009.',
        '[6] Center for Open Science, “OSF Preregistration,” https://osf.io/registries, accessed 2026.',
        '[7] NeuroTrace-DAG contributors, “Preregistration packet prereg-v1,” GitHub release, 2026. [Online]. Available: https://github.com/Oscar-man-shrestha/eval-validity-edu/releases/tag/prereg-v1',
    ]
    for ref in refs:
        p = add_para(doc, ref, size=9, space_after=2)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.left_indent = Cm(0.5)

    # footer note
    note = add_para(
        doc,
        "Status note for advisors: confirmatory H-rev scores are pending by protocol. "
        "Do not interpret Table II as empty results—the freeze holdout is locked until registration.",
        size=8,
        italic=True,
        space_before=12,
    )
    for run in note.runs:
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
